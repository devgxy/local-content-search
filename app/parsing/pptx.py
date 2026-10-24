"""PPTX 解析：文本框/表格/图表/图片 OCR/备注。

Step: STEP-03
Doc: docs/steps/STEP-03-解析层.md §5.4
"""

# =============================================================================
# 代码导读：PptxParser
# =============================================================================
#
# 职责：把 .pptx 转成 ParsedDocument；不做切片/向量化/写库。
#
# 输入：PPTX 路径 + Settings
# 输出：
#   ParsedDocument(
#       file_type="pptx",
#       title=第一页标题或文件名,
#       blocks=[heading/paragraph/table/ocr/note...],
#       meta={"slides": 页数},
#   )
#
# 主流程：
#
#   parse()
#     -> python-pptx Presentation(path)
#     -> 逐页（slide_no 从 1 开始）
#          -> 标题 -> heading
#          -> 背景图 OCR（可选）
#          -> 形状按 (top,left) 排序
#               -> 组合形状递归展开并排序
#               -> 文本框 -> 每个 paragraph 一个 block
#               -> 表格   -> 表头/行结构 -> table block
#               -> 图表   -> 类别/系列 -> table block
#               -> 图片   -> OCR / alt text
#               -> SmartArt/图形文本兜底
#          -> 备注 -> note
#          -> 备注图片 OCR
#     -> finalize_document
#
# 顺序：
#   标题 -> 背景图 -> 形状（top/left） -> 备注 -> 备注图片。
#   备注不在 shapes 流中，所以固定放在本页最后。
#
# P1：文本框按 paragraph 拆；保留 paragraph.level。
# P2：形状按 top/left 几何排序；同坐标保持原始 XML 顺序。
# P4：表格输出“表头/行”文本 + meta 中的 headers/rows。
# P3：图表提取标题、类别、系列、数值；SmartArt 从 diagram XML 兜底。
# P5：图片占位符/背景图/备注图/alt text 都本地处理。
#
# 示例 PPTX：
#
#   Slide 1:
#     Title: 系统架构
#     Text:  第一段 / 第二段
#     Table: 模块 | 说明\n解析 | PDF/DOCX
#     Chart: 月度请求量：1月=100, 2月=200
#     Image: 一张架构图，alt="系统架构图"
#     Notes: 备注：本页讲整体架构。
#
# 输出（节选）：
#     heading    "系统架构"
#     paragraph  "第一段"
#     paragraph  "第二段"
#     table      "表头：模块 | 说明\n行 1：解析 | PDF/DOCX"
#     table      "图表：月度请求量\n类别：1月 | 2月\n系列 请求量：100 | 200"
#     paragraph  "系统架构图"            # alt text
#     ocr        "架构图 OCR 文本"
#     note       "备注：本页讲整体架构。"
#
# 已知边界：
#   - heading_path 只有本页标题，没有标题层级；
#   - 表格整表一个 block；合并单元格已跳过 spanned cell；
#   - 图表只读取 python-pptx 可访问的类别/系列；
#   - SmartArt 依赖 diagram XML，复杂布局可能不完整；
#   - 只处理图片关系可访问的图片；背景图只读 slide 自身背景；
#   - 每页最多 20 张图片；同一图片重复出现不按 rId 去重（PPTX 无稳定 rId API）；
#   - 备注图片追加在备注文本之后；
#   - PPTX 没有 page 字段，只有 slide。
# =============================================================================

from __future__ import annotations

from pathlib import Path

from app.parsing.base import ParsedBlock, ParsedDocument, ParseError, finalize_document
from app.parsing.image_ocr import ocr_blob

# 单页 OCR 上限：背景图/形状图/备注图都计入，防止大 deck 拖垮任务。
MAX_IMAGES_PER_SLIDE = 20

# OOXML 关系属性：图片 r:embed；SmartArt 的 dm/lo/qs/cs。
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_R_DM = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}dm"
_R_LO = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}lo"
_R_QS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}qs"
_R_CS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}cs"


def _coord(value) -> int:  # noqa: ANN001
    """把 shape 的 top/left 坐标转成可排序整数。"""
    try:
        return int(value)
    except (TypeError, ValueError):
        # 某些占位符/组合形状可能没有显式坐标；放到 0 并保持稳定顺序。
        return 0


def _shape_key(shape) -> tuple[int, int]:  # noqa: ANN001
    """视觉阅读顺序键：先上后下，再左后右。"""
    return _coord(getattr(shape, "top", None)), _coord(getattr(shape, "left", None))


def _iter_shapes(shapes):  # noqa: ANN001
    """按几何位置递归遍历组合形状。

    Args:
        shapes: slide.shapes 或 group.shapes。

    Yields:
        非 GROUP 的 shape；组合形状按组合 bbox 排序后深度优先展开。
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    # enumerate 保留原始 XML 顺序，作为 top/left 相同时的稳定兜底。
    indexed = list(enumerate(shapes))
    indexed.sort(key=lambda item: (*_shape_key(item[1]), item[0]))
    for _, shape in indexed:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _fmt_value(value) -> str:  # noqa: ANN001
    """图表数值转文本：整数不带 .0。"""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _alt_text(shape) -> str:  # noqa: ANN001
    """读取 shape 的 alt text（cNvPr@descr）。"""
    try:
        nodes = shape._element.xpath(".//p:cNvPr")
        if nodes:
            return str(nodes[0].get("descr") or "").strip()
    except Exception:  # noqa: BLE001
        pass
    return ""


def _image_blob(shape):  # noqa: ANN001
    """尝试取图片字节；占位符图片也走这里。"""
    try:
        return shape.image.blob
    except Exception:  # noqa: BLE001
        return None


def _table_content(table) -> tuple[list[str], list[list[str]]]:
    """把 PPTX 表格转成 headers/rows，跳过合并单元格的 spanned 副本。

    Returns:
        ``(headers, rows)``；headers 是首个非空行，rows 不含表头。
    """
    grid: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            # 合并单元格的非起点会 is_spanned=True，避免重复文本。
            if getattr(cell, "is_spanned", False):
                continue
            cells.append(cell.text.strip())
        if any(cells):
            grid.append(cells)
    if not grid:
        return [], []
    return grid[0], grid[1:]


def _table_text(headers: list[str], rows: list[list[str]]) -> str:
    """表格结构化文本：首行表头 + 行号，便于切片层重复表头。"""
    lines = []
    if headers:
        lines.append("表头：" + " | ".join(headers))
    for index, row in enumerate(rows, start=1):
        lines.append(f"行 {index}：" + " | ".join(row))
    return "\n".join(lines)


def _chart_text(chart, fallback_name: str) -> str:  # noqa: ANN001
    """提取图表标题/类别/系列/数值，序列化为可检索文本。

    Args:
        chart: python-pptx Chart。
        fallback_name: 无标题时用 shape.name 兜底。

    Returns:
        多行文本，交给 table 类型切片。
    """
    title = ""
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text.strip()
    except Exception:  # noqa: BLE001
        title = ""
    lines = [f"图表：{title or fallback_name}"]
    try:
        for plot in chart.plots:
            categories = []
            try:
                categories = [str(item) for item in (plot.categories or [])]
            except Exception:  # noqa: BLE001
                categories = []
            if categories:
                lines.append("类别：" + " | ".join(categories))
            for series in plot.series:
                name = str(getattr(series, "name", "") or "系列")
                values = []
                try:
                    values = [_fmt_value(value) for value in series.values]
                except Exception:  # noqa: BLE001
                    values = []
                if values:
                    lines.append(f"系列 {name}：" + " | ".join(values))
    except Exception:  # noqa: BLE001
        pass
    return "\n".join(lines)


def _diagram_texts(shape) -> list[str]:  # noqa: ANN001
    """从 SmartArt diagram XML 提取文本，作为无高层 API 的兜底。

    python-pptx 不暴露 SmartArt 文本，因此沿 dgm:relIds 找 data part，
    再取其中 a:t 文本；失败时退回 shape XML 的 a:t。
    """
    texts: list[str] = []

    def add(value: str) -> None:
        value = value.strip()
        if value and value not in texts:
            texts.append(value)

    try:
        from lxml import etree

        rel_ids = shape._element.xpath(".//dgm:relIds")
        for node in rel_ids:
            for attr in (_R_DM, _R_LO, _R_QS, _R_CS):
                r_id = node.get(attr)
                if not r_id:
                    continue
                rel = shape.part.rels.get(r_id)
                if rel is None:
                    continue
                try:
                    root = etree.fromstring(rel.target_part.blob)
                    for text_node in root.xpath(
                        ".//a:t",
                        namespaces={"a": "http://schemas.openxmlformats.org/drawingml/2006/main"},
                    ):
                        add(text_node.text or "")
                except Exception:  # noqa: BLE001
                    continue
    except Exception:  # noqa: BLE001
        pass
    if not texts:
        try:
            for text_node in shape._element.xpath(".//a:t"):
                add(text_node.text or "")
        except Exception:  # noqa: BLE001
            pass
    return texts


class PptxParser:
    """PPTX 解析器：形状/表格/图表/图片 OCR/备注。"""

    extensions = (".pptx",)

    def parse(self, path: Path, settings) -> ParsedDocument:  # noqa: ANN001
        """解析 PPTX 文件。

        Args:
            path: PPTX 绝对路径。
            settings: OCR/置信度/最小像素等配置。

        Returns:
            ParsedDocument；表格/图表已结构化文本化，图片按形状顺序 OCR。
        """
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ParseError(
                "BROKEN_FILE", "未安装 python-pptx", "pip install python-pptx"
            ) from exc

        try:
            presentation = Presentation(str(path))
        except Exception as exc:  # noqa: BLE001
            raise ParseError("BROKEN_FILE", f"无法打开 PPTX：{exc}") from exc

        warnings: list[str] = []
        blocks: list[ParsedBlock] = []
        title = ""

        for slide_no, slide in enumerate(presentation.slides, start=1):
            # 每页独立 section；slide_no 是 1-based，后续 block 都带这个值。
            title_shape = slide.shapes.title
            slide_title = title_shape.text.strip() if title_shape is not None else ""
            # 文档 title 取第一页非空标题；每页 heading_path 只有本页标题。
            if slide_title and not title:
                title = slide_title
            heading_path = (slide_title,) if slide_title else ()
            if slide_title:
                blocks.append(
                    ParsedBlock(
                        text=slide_title,
                        kind="heading",
                        slide=slide_no,
                        heading_path=heading_path,
                    )
                )

            image_count = 0
            limit_warned = False
            ocr_disabled_warned = False

            def add_image(
                blob: bytes,
                alt: str,
                source: str,
                *,
                slide_no: int = slide_no,
                heading_path: tuple[str, ...] = heading_path,
            ) -> None:
                """OCR 一张图片，并在当前位置追加 alt/OCR block。

                slide_no/heading_path 用默认参数绑定当前循环值，避免 closure
                在后续幻灯片迭代时读到新的值。
                """
                nonlocal image_count, limit_warned, ocr_disabled_warned
                if not settings.ocr_enabled:
                    if not ocr_disabled_warned:
                        warnings.append("OCR 已关闭，跳过 PPTX 图片")
                        ocr_disabled_warned = True
                    return
                if image_count >= MAX_IMAGES_PER_SLIDE:
                    if not limit_warned:
                        warnings.append(
                            f"第 {slide_no} 页图片超过 {MAX_IMAGES_PER_SLIDE} 张，其余跳过"
                        )
                        limit_warned = True
                    return
                image_count += 1
                try:
                    text, image_warnings = ocr_blob(
                        blob,
                        min_confidence=settings.ocr_min_confidence,
                        min_px=settings.ocr_min_image_px,
                    )
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"第 {slide_no} 页图片 OCR 失败：{exc}")
                    return
                warnings.extend(f"第 {slide_no} 页：{w}" for w in image_warnings)
                # P5：alt text 作为可检索的描述，放在 OCR 文本前。
                if alt and alt not in text:
                    blocks.append(
                        ParsedBlock(
                            text=alt,
                            kind="paragraph",
                            slide=slide_no,
                            heading_path=heading_path,
                            meta={"source": f"{source}_alt"},
                        )
                    )
                if text.strip():
                    blocks.append(
                        ParsedBlock(
                            text=text,
                            kind="ocr",
                            slide=slide_no,
                            heading_path=heading_path,
                            meta={"source": source, "alt_text": alt},
                        )
                    )

            # P5：slide 自身背景图（不包含 layout/master 背景）。
            try:
                for blip in slide._element.xpath(".//p:bg//a:blip"):
                    r_id = blip.get(_R_EMBED)
                    part = slide.part.related_parts.get(r_id) if r_id else None
                    if part is not None:
                        add_image(part.blob, "", "slide_background")
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"第 {slide_no} 页背景图读取失败：{exc}")

            # P2：按几何位置排序；组合形状内部也按位置排序。
            for shape in _iter_shapes(slide.shapes):
                # 标题已经输出过 heading block，这里跳过，避免重复。
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue

                if getattr(shape, "has_text_frame", False):
                    # P1：一个 text frame 按 paragraph 拆；每个段落独立成 block。
                    for para_index, paragraph in enumerate(shape.text_frame.paragraphs):
                        text = paragraph.text.strip()
                        if not text:
                            continue
                        blocks.append(
                            ParsedBlock(
                                text=text,
                                kind="paragraph",
                                slide=slide_no,
                                heading_path=heading_path,
                                meta={
                                    "paragraph_index": para_index,
                                    "paragraph_level": paragraph.level,
                                },
                            )
                        )
                    continue

                if getattr(shape, "has_table", False):
                    # P4：表头/行结构，文本化后交给切片层按行处理。
                    headers, rows = _table_content(shape.table)
                    if headers or rows:
                        blocks.append(
                            ParsedBlock(
                                text=_table_text(headers, rows),
                                kind="table",
                                slide=slide_no,
                                heading_path=heading_path,
                                meta={
                                    "source": "pptx_table",
                                    "table_headers": headers,
                                    "table_rows": rows,
                                },
                            )
                        )
                    continue

                if getattr(shape, "has_chart", False):
                    # P3：图表转“类别/系列/数值”文本，kind=table 便于按行切。
                    blocks.append(
                        ParsedBlock(
                            text=_chart_text(shape.chart, shape.name),
                            kind="table",
                            slide=slide_no,
                            heading_path=heading_path,
                            meta={"source": "pptx_chart"},
                        )
                    )
                    continue

                # P5：不只看 PICTURE；图片占位符也暴露 .image。
                blob = _image_blob(shape)
                if blob is not None:
                    add_image(blob, _alt_text(shape), "pptx_image")
                    continue

                # P3：SmartArt/图形框兜底，从 diagram XML 取文本。
                # 不依赖特定 shape_type：不同 python-pptx 版本的枚举不一致。
                diagram_texts = _diagram_texts(shape)
                if diagram_texts:
                    blocks.append(
                        ParsedBlock(
                            text="\n".join(diagram_texts),
                            kind="paragraph",
                            slide=slide_no,
                            heading_path=heading_path,
                            meta={"source": "pptx_diagram"},
                        )
                    )

            # 备注不在 shapes 流中，固定追加到本页最后。
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    blocks.append(
                        ParsedBlock(
                            text=notes,
                            kind="note",
                            slide=slide_no,
                            heading_path=heading_path,
                            meta={"source": "pptx_notes"},
                        )
                    )
                # P5：备注页图片 OCR（放在备注文本之后）。
                try:
                    for note_shape in _iter_shapes(slide.notes_slide.shapes):
                        blob = _image_blob(note_shape)
                        if blob is not None:
                            add_image(blob, _alt_text(note_shape), "notes_image")
                except Exception as exc:  # noqa: BLE001
                    warnings.append(f"第 {slide_no} 页备注图片读取失败：{exc}")

        # finalize 负责丢空 block、计算 char 偏移、生成 doc.text。
        return finalize_document(
            ParsedDocument(
                path=path,
                file_type="pptx",
                title=title or path.stem,
                blocks=blocks,
                warnings=warnings,
                meta={"slides": len(presentation.slides)},
            )
        )
