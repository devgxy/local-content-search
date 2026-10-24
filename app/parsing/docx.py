"""DOCX 解析：保序段落/标题/表格/图片 OCR。

主流程：
1. python-docx 打开文档；
2. iter_inner_content() 按文档顺序遍历 paragraph/table；
3. 段落内按 run 顺序处理文本和内嵌图片；
4. 表格单元格内的图片递归处理；
5. 图片 OCR block 插入到当前位置，并使用当前 heading_path；
6. finalize_document 计算偏移/生成全文。

Step: STEP-03
Doc: docs/steps/STEP-03-解析层.md §5.3
"""

# =============================================================================
# 代码导读：DocxParser
# =============================================================================
#
# 职责：把 .docx 转成 ParsedDocument；不做切片/向量化/写库。
#
# 输入：DOCX 路径 + Settings
# 输出：
#   ParsedDocument(
#       file_type="docx",
#       title=第一个 Heading 1 或文件名,
#       blocks=[heading/paragraph/table/ocr...],
#       warnings=[...],
#   )
#
# 与旧实现的关键区别：
#   旧实现：先遍历全部段落/表格，再从 document.part.rels 统一 OCR 图片，
#           导致所有 OCR block 追加到末尾，且 heading_path 是最后一个标题。
#   新实现：遍历段落时按 run 顺序找图片，图片 OCR block 插入当前位置，
#           使用图片出现时的 heading_stack。
#
# 文档顺序来源：
#   document.iter_inner_content() -> Paragraph/Table（顶层）
#   Paragraph.runs -> Run（段落内顺序）
#   Run._element.xpath(".//a:blip") -> DrawingML 图片（r:embed = rId）
#
# 示例 DOCX：
#
#   Heading 1: 错误码
#   Paragraph: 下表列出错误码。
#   Image: 截图 A
#   Heading 2: 项目经验
#   Paragraph: 项目内容。
#   Image: 截图 B
#
# 新实现输出顺序：
#
#   heading    "错误码"                 heading_path=["错误码"]
#   paragraph  "下表列出错误码。"        heading_path=["错误码"]
#   ocr        "截图 A OCR 文本"         heading_path=["错误码"]
#   heading    "项目经验"               heading_path=["错误码","项目经验"]
#   paragraph  "项目内容。"             heading_path=["错误码","项目经验"]
#   ocr        "截图 B OCR 文本"         heading_path=["错误码","项目经验"]
#
# 已知限制：
#   - 只处理 DrawingML 的 a:blip；旧式 VML v:imagedata 暂不处理；
#   - 段落内“文本 + 图片 + 文本”在同一个 run 时，run.text 会整体放在
#     图片之前，这是 python-docx run API 的近似；
#   - 浮动图片（wp:anchor）按锚点所在段落定位，不等于视觉位置；
#   - 表格图片追加在 table block 之后，不还原单元格内的精确位置；
#   - 嵌套表格、文本框、页眉页脚、脚注/批注不处理；
#   - 合并单元格在 row.cells 中可能重复，已按 _tc 去重图片；
#   - 同一 rId 的图片只 OCR 一次，避免 logo 重复产生重复 chunk；
#   - DOCX 没有可靠页码，block 的 page 固定为 0；
#   - OCR block 的 char_start/char_end 由 finalize 按当前顺序计算。
# =============================================================================

from __future__ import annotations

import re
from pathlib import Path

from app.parsing.base import (
    ParsedBlock,
    ParsedDocument,
    ParseError,
    finalize_document,
    normalize_text,
)
from app.parsing.image_ocr import ocr_blob

_HEADING_RE = re.compile(r"(?:Heading|标题)\s*(\d)", re.IGNORECASE)
MAX_IMAGES = 50

# DrawingML inline/anchor 图片的关系属性。
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _heading_level(paragraph) -> int:  # noqa: ANN001
    """根据样式名推断 Markdown 标题级别。"""
    # 同时看直接样式和 base_style，兼容“标题 1”继承“Heading 1”的模板。
    names = []
    if paragraph.style is not None:
        names.append(paragraph.style.name or "")
        if paragraph.style.base_style is not None:
            names.append(paragraph.style.base_style.name or "")
    for name in names:
        match = _HEADING_RE.search(name)
        if match:
            return int(match.group(1))
    return 0


class DocxParser:
    """DOCX 解析器：段落/表格/图片 OCR，按文档顺序输出 block。"""

    extensions = (".docx",)

    def parse(self, path: Path, settings) -> ParsedDocument:  # noqa: ANN001
        """解析 DOCX 文件。

        Args:
            path: DOCX 绝对路径。
            settings: OCR 开关/置信度/最小像素等配置。

        Returns:
            ParsedDocument；图片 OCR 按文档位置插入。

        Raises:
            ParseError: 缺少 python-docx / 文件损坏。
        """
        # 延迟导入，只有解析 DOCX 才依赖 python-docx。
        try:
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as exc:
            raise ParseError(
                "BROKEN_FILE", "未安装 python-docx", "pip install python-docx"
            ) from exc

        try:
            document = Document(str(path))
        except Exception as exc:  # noqa: BLE001
            raise ParseError("BROKEN_FILE", f"无法打开 DOCX：{exc}") from exc

        warnings: list[str] = []
        blocks: list[ParsedBlock] = []
        heading_stack: list[str] = []
        title = ""
        image_count = 0
        seen_images: set[str] = set()
        ocr_disabled_warned = False
        limit_warned = False

        def flush_text(buffer: list[str], heading: tuple[str, ...]) -> None:
            """把 run 文本缓冲刷成一个 paragraph block。"""
            body = normalize_text("".join(buffer))
            if body:
                blocks.append(
                    ParsedBlock(
                        text=body,
                        kind="paragraph",
                        heading_path=heading,
                    )
                )
            buffer.clear()

        def add_image(r_id: str, heading: tuple[str, ...]) -> None:
            """按 rId 取图片并 OCR，在当前文档位置追加 ocr block。

            Args:
                r_id: DOCX 关系 ID（r:embed）。
                heading: 图片出现时的标题路径快照。
            """
            nonlocal image_count, ocr_disabled_warned, limit_warned

            if not settings.ocr_enabled:
                if not ocr_disabled_warned:
                    warnings.append("OCR 已关闭，跳过内嵌图片")
                    ocr_disabled_warned = True
                return
            # 同一关系只处理一次；同一 logo 重复引用不会产生重复 chunk。
            if r_id in seen_images:
                return
            if image_count >= MAX_IMAGES:
                if not limit_warned:
                    warnings.append(f"图片超过 {MAX_IMAGES} 张，其余跳过")
                    limit_warned = True
                return
            part = document.part.related_parts.get(r_id)
            if part is None:
                warnings.append(f"图片关系不存在：{r_id}")
                return
            try:
                blob = part.blob
            except Exception as exc:  # noqa: BLE001
                warnings.append(f"读取内嵌图片失败：{exc}")
                return
            seen_images.add(r_id)
            image_count += 1
            text, image_warnings = ocr_blob(
                blob,
                min_confidence=settings.ocr_min_confidence,
                min_px=settings.ocr_min_image_px,
            )
            warnings.extend(image_warnings)
            if text.strip():
                blocks.append(
                    ParsedBlock(
                        text=text,
                        kind="ocr",
                        heading_path=heading,
                        meta={"source": "docx_image", "r_id": r_id},
                    )
                )

        def paragraph_images(paragraph) -> list[object]:  # noqa: ANN001
            """返回段落内所有 DrawingML blip 元素。"""
            try:
                return list(paragraph._element.xpath(".//a:blip"))
            except Exception:  # noqa: BLE001
                return []

        def process_paragraph(paragraph, *, include_text: bool = True) -> None:  # noqa: ANN001
            """处理一个段落：标题/文本/段落内图片。

            Args:
                paragraph: python-docx Paragraph。
                include_text: False 时只处理图片（表格单元格用），不重复正文。
            """
            nonlocal heading_stack, title
            text = paragraph.text.strip()
            level = _heading_level(paragraph)
            has_images = bool(paragraph_images(paragraph))

            if include_text and level:
                # 标题：更新标题栈并输出 heading block；标题文字不进入正文。
                heading_stack = heading_stack[: level - 1] + [text]
                if level == 1 and not title:
                    title = text
                if text:
                    blocks.append(
                        ParsedBlock(
                            text=text,
                            kind="heading",
                            heading_path=tuple(heading_stack),
                        )
                    )
                # 标题段落里的图片也要处理，但不再输出段落文本。
                include_text = False
            elif include_text and not has_images:
                # 没有图片：保留旧行为，直接用 paragraph.text。
                if text:
                    blocks.append(
                        ParsedBlock(
                            text=text,
                            kind="paragraph",
                            heading_path=tuple(heading_stack),
                        )
                    )
                return
            elif include_text and text:
                # 有图片：按下述 run 顺序拆分文本/图片。
                pass
            elif not include_text and not has_images:
                return

            # 按 run 顺序累积文本；遇到图片先 flush 文本，再插入 OCR block。
            buffer: list[str] = []
            current_heading = tuple(heading_stack)
            for run in paragraph.runs:
                if include_text and run.text:
                    buffer.append(run.text)
                for blip in run._element.xpath(".//a:blip"):
                    r_id = blip.get(_R_EMBED)
                    if not r_id:
                        continue
                    if include_text:
                        flush_text(buffer, current_heading)
                    add_image(r_id, current_heading)
            if include_text:
                flush_text(buffer, current_heading)

        def process_table(table) -> None:  # noqa: ANN001
            """输出 table block，并处理单元格内的图片。

            table 文本会先整体输出；单元格图片随后追加到 table block 之后。
            这是为了不破坏表格行的文本结构。
            """
            rows = []
            for row in table.rows:
                # 合并单元格在 row.cells 中可能重复，文本会重复出现。
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                blocks.append(
                    ParsedBlock(
                        text="\n".join(rows),
                        kind="table",
                        heading_path=tuple(heading_stack),
                    )
                )

            # 单元格图片：只处理图片，不再输出 cell 文本（已进 table block）。
            seen_cells: set[int] = set()
            for row in table.rows:
                for cell in row.cells:
                    cell_key = id(cell._tc)
                    if cell_key in seen_cells:
                        continue
                    seen_cells.add(cell_key)
                    for paragraph in cell.paragraphs:
                        process_paragraph(paragraph, include_text=False)

        # iter_inner_content() 保持顶层段落/表格的文档顺序。
        for item in document.iter_inner_content():
            if isinstance(item, Paragraph):
                process_paragraph(item)
            elif isinstance(item, Table):
                process_table(item)

        # finalize 负责丢空 block、计算 char 偏移、生成 doc.text。
        return finalize_document(
            ParsedDocument(
                path=path,
                file_type="docx",
                title=title or path.stem,
                blocks=blocks,
                warnings=warnings,
            )
        )
