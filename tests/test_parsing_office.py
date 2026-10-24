"""STEP-03 DOCX/PPTX 解析测试。"""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from app.config import PROJECT_ROOT, Settings
from app.parsing.registry import get_parser


def test_docx(fixtures_dir: Path, tmp_settings: Settings) -> None:
    path = fixtures_dir / "sample.docx"
    doc = get_parser(path).parse(path, tmp_settings)
    assert doc.file_type == "docx"
    assert doc.title == "混合检索系统设计"
    assert any(block.kind == "table" and "模块 | 说明" in block.text for block in doc.blocks)
    assert any(block.kind == "ocr" and "混合检索" in block.text for block in doc.blocks)
    assert "第二段" in doc.text


def test_pptx(fixtures_dir: Path, tmp_settings: Settings) -> None:
    path = fixtures_dir / "sample.pptx"
    doc = get_parser(path).parse(path, tmp_settings)
    assert doc.file_type == "pptx"
    assert doc.title == "混合检索系统"
    assert any(block.kind == "note" and "备注" in block.text for block in doc.blocks)
    assert any(block.slide == 1 for block in doc.blocks)
    assert any(block.kind == "ocr" and "混合检索" in block.text for block in doc.blocks)


def test_office_ocr_disabled(fixtures_dir: Path, tmp_settings: Settings) -> None:
    settings = replace(tmp_settings, ocr_enabled=False)
    for name in ("sample.docx", "sample.pptx"):
        path = fixtures_dir / name
        doc = get_parser(path).parse(path, settings)
        assert not any(block.kind == "ocr" for block in doc.blocks)


def test_docx_image_order_and_heading(tmp_path: Path, fixtures_dir: Path, tmp_settings: Settings) -> None:
    from docx import Document
    from docx.shared import Inches

    path = tmp_path / "ordered.docx"
    document = Document()
    document.add_heading("错误码", level=1)
    document.add_paragraph("下表列出错误码。")
    document.add_picture(str(fixtures_dir / "sample.png"), width=Inches(2))
    document.add_heading("项目经验", level=2)
    document.add_paragraph("项目内容。")
    second_image = PROJECT_ROOT / "docs" / "assets" / "ui-search.png"
    document.add_picture(str(second_image), width=Inches(2))
    document.save(str(path))

    parsed = get_parser(path).parse(path, tmp_settings)
    kinds = [block.kind for block in parsed.blocks]
    assert kinds == ["heading", "paragraph", "ocr", "heading", "paragraph", "ocr"]
    assert parsed.blocks[2].heading_path == ("错误码",)
    assert parsed.blocks[5].heading_path == ("错误码", "项目经验")


def test_docx_image_before_first_heading(tmp_path: Path, fixtures_dir: Path, tmp_settings: Settings) -> None:
    from docx import Document
    from docx.shared import Inches

    path = tmp_path / "before-heading.docx"
    document = Document()
    document.add_picture(str(fixtures_dir / "sample.png"), width=Inches(2))
    document.add_heading("错误码", level=1)
    document.save(str(path))

    parsed = get_parser(path).parse(path, tmp_settings)
    assert parsed.blocks[0].kind == "ocr"
    assert parsed.blocks[0].heading_path == ()


def test_docx_table_cell_image(tmp_path: Path, fixtures_dir: Path, tmp_settings: Settings) -> None:
    from docx import Document
    from docx.shared import Inches

    path = tmp_path / "table-image.docx"
    document = Document()
    document.add_heading("表格", level=1)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "说明"
    table.cell(0, 0).add_paragraph().add_run().add_picture(
        str(fixtures_dir / "sample.png"), width=Inches(1)
    )
    document.save(str(path))

    parsed = get_parser(path).parse(path, tmp_settings)
    kinds = [block.kind for block in parsed.blocks]
    assert "table" in kinds
    assert "ocr" in kinds
    ocr = next(block for block in parsed.blocks if block.kind == "ocr")
    assert ocr.heading_path == ("表格",)


def test_pptx_paragraphs_split_and_levels(tmp_path: Path, tmp_settings: Settings) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "paragraphs.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "段落测试"
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    frame = box.text_frame
    frame.text = "第一段"
    second = frame.add_paragraph()
    second.text = "第二段"
    second.level = 1
    presentation.save(str(path))

    parsed = get_parser(path).parse(path, tmp_settings)
    paragraphs = [block for block in parsed.blocks if block.kind == "paragraph"]
    assert [block.text for block in paragraphs] == ["第一段", "第二段"]
    assert [block.meta["paragraph_level"] for block in paragraphs] == [0, 1]
    assert all(block.heading_path == ("段落测试",) for block in paragraphs)


def test_pptx_geometry_reading_order(tmp_path: Path, tmp_settings: Settings) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "geometry.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "顺序测试"
    # 先插入底部，再插入顶部；视觉顺序应为顶部 -> 底部。
    bottom = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(0.5))
    bottom.text_frame.text = "底部"
    top = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(3), Inches(0.5))
    top.text_frame.text = "顶部"
    presentation.save(str(path))

    parsed = get_parser(path).parse(path, tmp_settings)
    paragraphs = [block.text for block in parsed.blocks if block.kind == "paragraph"]
    assert paragraphs == ["顶部", "底部"]


def test_pptx_table_structure(tmp_path: Path, tmp_settings: Settings) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "table-structure.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "表格结构"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1.5), Inches(5), Inches(1.5)).table
    table.cell(0, 0).text = "模块"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "解析"
    table.cell(1, 1).text = "PDF/DOCX"
    presentation.save(str(path))

    parsed = get_parser(path).parse(path, tmp_settings)
    table_block = next(block for block in parsed.blocks if block.kind == "table")
    assert table_block.meta["table_headers"] == ["模块", "说明"]
    assert table_block.meta["table_rows"] == [["解析", "PDF/DOCX"]]
    assert table_block.text.startswith("表头：模块 | 说明")


def test_pptx_chart_text(tmp_path: Path, tmp_settings: Settings) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    path = tmp_path / "chart.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "图表页"
    data = CategoryChartData()
    data.categories = ["1月", "2月"]
    data.add_series("请求量", (100, 200))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1.5),
        Inches(6),
        Inches(4),
        data,
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "月度请求量"
    presentation.save(str(path))

    parsed = get_parser(path).parse(path, tmp_settings)
    chart_block = next(
        block for block in parsed.blocks if block.meta.get("source") == "pptx_chart"
    )
    assert "图表：月度请求量" in chart_block.text
    assert "类别：1月 | 2月" in chart_block.text
    assert "系列 请求量：100 | 200" in chart_block.text


def test_pptx_picture_placeholder_ocr(tmp_path: Path, fixtures_dir: Path, tmp_settings: Settings) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER

    path = tmp_path / "placeholder.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[8])
    slide.shapes.title.text = "图片占位符"
    picture = next(
        shape
        for shape in slide.placeholders
        if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    )
    picture = picture.insert_picture(str(fixtures_dir / "sample.png"))
    # P5：alt text 也要成为可检索文本。
    picture._element.xpath(".//p:cNvPr")[0].set("descr", "系统架构图")
    presentation.save(str(path))

    parsed = get_parser(path).parse(path, tmp_settings)
    assert any(
        block.kind == "ocr" and "混合检索" in block.text
        for block in parsed.blocks
    )
    assert any(
        block.meta.get("source") == "pptx_image_alt" and block.text == "系统架构图"
        for block in parsed.blocks
    )
