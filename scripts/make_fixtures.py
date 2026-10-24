"""生成解析层测试 fixtures。

Step: STEP-03
Doc: docs/steps/STEP-03-解析层.md §8.1
"""

from __future__ import annotations

import argparse
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches as PptxInches

CJK_FONT = "/System/Library/Fonts/Hiragino Sans GB.ttc"
PDF_FONT = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"


def make_image(path: Path) -> None:
    image = Image.new("RGB", (1500, 480), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(CJK_FONT, 46, index=0)
    body_font = ImageFont.truetype(CJK_FONT, 34, index=0)
    draw.text((40, 40), "混合检索系统：支持 PDF、DOCX、PPTX 与图片 OCR", font=title_font, fill="black")
    draw.text((40, 140), "错误码 E1024：连接超时，请检查代理配置。", font=body_font, fill="black")
    draw.text((40, 210), "BGE 重排模型可以在 Mac 本地离线运行。", font=body_font, fill="black")
    image.save(path)


def make_text(out: Path) -> None:
    (out / "tiny.md").write_text(
        "# 混合检索系统\n\n"
        "第一段：向量召回与关键词匹配。\n\n"
        "## 检索策略\n\n"
        "第二段：RRF 融合后使用 bge-reranker-large 重排。\n\n"
        "```python\nprint('hello')\n```\n\n"
        "| 模块 | 说明 |\n|---|---|\n| 解析 | PDF/DOCX/PPTX |\n",
        encoding="utf-8",
    )
    (out / "gb18030.txt").write_text(
        "中文编码测试：混合检索系统。\n\n第二段：错误码 E1024。",
        encoding="gb18030",
    )


def make_pdf(out: Path, image_path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_font(fontname="cjk", fontfile=PDF_FONT)
    page.insert_text((72, 100), "混合检索系统设计文档", fontname="cjk", fontsize=20)
    page.insert_text((72, 140), "支持 PDF、DOCX、PPTX 与图片 OCR。", fontname="cjk", fontsize=12)
    page.insert_text((72, 180), "错误码 E1024：连接超时。", fontname="cjk", fontsize=12)
    doc.save(str(out / "sample.pdf"))
    doc.close()

    scanned = pymupdf.open()
    page = scanned.new_page(width=595, height=420)
    page.insert_image(pymupdf.Rect(0, 0, 595, 420), filename=str(image_path))
    scanned.save(str(out / "scanned.pdf"))
    scanned.close()

    broken = out / "broken.pdf"
    broken.write_bytes((out / "sample.pdf").read_bytes()[:800])


def make_docx(out: Path, image_path: Path) -> None:
    document = Document()
    document.add_heading("混合检索系统设计", level=1)
    document.add_paragraph("第一段：支持 PDF、DOCX、PPTX 与图片 OCR。")
    document.add_heading("检索策略", level=2)
    document.add_paragraph("第二段：向量召回 + BM25 + RRF + 重排。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "模块"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "解析"
    table.cell(1, 1).text = "PDF/DOCX/PPTX"
    document.add_picture(str(image_path), width=Inches(5))
    document.save(str(out / "sample.docx"))


def make_pptx(out: Path, image_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "混合检索系统"
    slide.placeholders[1].text = "向量召回 + 关键词匹配 + RRF + 重排"
    slide.notes_slide.notes_text_frame.text = "备注：本页讲整体架构。"
    slide.shapes.add_picture(str(image_path), PptxInches(1), PptxInches(3), width=PptxInches(4))
    presentation.save(str(out / "sample.pptx"))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="tests/fixtures")
    args = parser.parse_args()

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    image_path = out / "sample.png"
    make_image(image_path)
    make_text(out)
    make_pdf(out, image_path)
    make_docx(out, image_path)
    make_pptx(out, image_path)
    print(f"fixtures -> {out.resolve()}")


if __name__ == "__main__":
    main()
